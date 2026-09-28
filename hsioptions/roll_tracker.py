#!/usr/bin/env python3
"""
期指/期權轉倉位自動追蹤表
用法: python3 roll_tracker.py OCT-26 2026-09-14 2026-09-23
輸出: ~/workspace/your_files/hsioptions/轉倉位追蹤_OCT-26.xlsx
版面仿照用戶原有的手動 Google Sheet: A欄記號(★▲✓), 每個交易日兩欄 [未平倉][漲跌],
首欄為期間總漲跌, 尾欄為最近2天漲跌, 頂部列出恆指收市位。
"""
import csv, os, sys, subprocess
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA = os.path.expanduser('~/workspace/hsioptions/data')
OUTDIR = os.path.expanduser('~/workspace/your_files/hsioptions')
BASE = 'https://7desl.com/hkex/data'
UA = {'User-Agent': 'Mozilla/5.0'}

YELLOW = PatternFill('solid', fgColor='FFFF00')
PINK = PatternFill('solid', fgColor='F4B8C1')
GREEN = PatternFill('solid', fgColor='B6E3B6')
HDR_FILL = PatternFill('solid', fgColor='EBF7FE')
SEC_FILL = PatternFill('solid', fgColor='4472C4')
THIN = Border(*(Side(style='thin', color='D9D9D9') for _ in range(4)))
CENTER = Alignment(horizontal='center', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')


def daterange(a, b):
    d = a
    while d <= b:
        yield d
        d += timedelta(days=1)


def dl(url, path):
    if os.path.exists(path) and os.path.getsize(path) > 0:
        return True
    r = subprocess.run(['curl', '-s', '-o', path, '-w', '%{http_code}', '-H',
                        f"User-Agent: {UA['User-Agent']}", url], capture_output=True, text=True)
    if r.stdout.strip() != '200' or os.path.getsize(path) == 0:
        if os.path.exists(path):
            os.remove(path)
        return False
    return True


def load_day(day, month):
    """回傳 {strike: dict}；None 表示該日無數據"""
    p = os.path.join(DATA, f'oi-{month}-{day}.csv')
    if not dl(f'{BASE}/{day}/hsi-options-months-{month}.csv', p):
        return None
    out = {}
    with open(p, encoding='utf-8-sig') as fh:
        rows = list(csv.reader(fh))
    for r in rows[1:]:
        if len(r) < 25 or not r[12].strip().lstrip('-').isdigit():
            continue
        s = int(r[12])
        out[s] = dict(c_oi=int(r[11] or 0), c_chg=int(r[9] or 0),
                      p_oi=int(r[13] or 0), p_chg=int(r[15] or 0))
    return out


def load_hsi(day):
    p = os.path.join(DATA, f'hsi-index-{day}.csv')
    if not dl(f'{BASE}/{day}/data-hsi-index.csv', p):
        return ''
    with open(p, encoding='utf-8-sig') as fh:
        rows = list(csv.reader(fh))
    if len(rows) > 1 and len(rows[1]) > 4:
        try:
            return f"{float(rows[1][2]):,.0f}"
        except ValueError:
            return ''
    return ''


def build(month, start, end):
    days, series = [], []
    for d in daterange(start, end):
        ds = d.isoformat()
        s = load_day(ds, month)
        if s is not None:
            days.append(ds)
            series.append(s)
    if not days:
        print('no data'); return
    hsi = [load_hsi(d) for d in days]
    strikes = sorted(set().union(*[set(s) for s in series]))
    # 只保留有意義的行使價: 期間總OI>0
    strikes = [s for s in strikes if sum((series[-1].get(s, {}).get('c_oi', 0) +
                                          series[-1].get(s, {}).get('p_oi', 0)) for _ in [0]) > 0
               or any(series[i].get(s, {}).get('c_oi', 0) + series[i].get(s, {}).get('p_oi', 0) > 0 for i in range(len(days)))]

    def diffs(s, key, chgkey):
        vals, chgs = [], []
        prev = None
        for i, ser in enumerate(series):
            v = ser.get(s, {}).get(key, 0)
            vals.append(v)
            chgs.append(ser.get(s, {}).get(chgkey, 0) if i == 0 else v - prev)
            prev = v
        return vals, chgs

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'{month}轉倉位'
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    n = len(days)
    # 欄: A=記號(★▲✓, 用戶手填), B=行使價, C=期間漲跌, 然後每日2欄[未平倉][漲跌], 最後2天漲跌1欄
    total_cols = 3 + 2 * n + 1

    def section(title, key, chgkey, r0, put=True):
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=total_cols)
        c = ws.cell(r0, 1, title)
        c.fill = SEC_FILL; c.font = Font(bold=True, color='FFFFFF', size=12); c.alignment = CENTER
        # 日期列
        ws.cell(r0 + 1, 2, '行使價').alignment = CENTER
        c2 = ws.cell(r0 + 1, 3, f'{days[0][5:]} 至 {days[-1][5:]}\n期間漲跌')
        c2.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, d in enumerate(days):
            c1 = 4 + 2 * i
            ws.merge_cells(start_row=r0 + 1, start_column=c1, end_row=r0 + 1, end_column=c1 + 1)
            ws.cell(r0 + 1, c1, d[5:].replace('-', '/')).alignment = CENTER
            ws.cell(r0 + 2, c1, '未平倉').alignment = CENTER
            ws.cell(r0 + 2, c1 + 1, hsi[i]).alignment = CENTER
        ws.cell(r0 + 1, 4 + 2 * n, '2天漲跌').alignment = CENTER
        for rr in range(r0 + 1, r0 + 3):
            for cc in range(1, total_cols + 1):
                cell = ws.cell(rr, cc)
                cell.fill = HDR_FILL; cell.border = THIN
                cell.font = Font(bold=True, size=10)
        # 數據列
        r = r0 + 3
        for s in strikes:
            vals, chgs = diffs(s, key, chgkey)
            tot = sum(chgs)
            last2 = sum(chgs[-2:])
            ws.cell(r, 2, s).alignment = CENTER
            ws.cell(r, 3, tot).alignment = RIGHT
            if abs(tot) >= 500:
                ws.cell(r, 3).fill = PINK if tot > 0 else GREEN
            maxv = max(vals) if vals else 0
            for i, (v, ch) in enumerate(zip(vals, chgs)):
                c1 = 4 + 2 * i
                vc = ws.cell(r, c1, v); vc.alignment = RIGHT
                if v == maxv and v > 0:
                    vc.fill = YELLOW
                cc = ws.cell(r, c1 + 1, ch); cc.alignment = RIGHT
                if abs(ch) >= 200:
                    cc.fill = PINK if ch > 0 else GREEN
            ws.cell(r, 4 + 2 * n, last2).alignment = RIGHT
            for cc in range(1, total_cols + 1):
                ws.cell(r, cc).border = THIN
                ws.cell(r, cc).font = Font(size=10)
            r += 1
        return r + 2

    r = 1
    r = section('認沽', 'p_oi', 'p_chg', r, put=True)
    section('認購', 'c_oi', 'c_chg', r, put=False)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 12
    for i in range(4, total_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 10
    ws.freeze_panes = 'D4'

    os.makedirs(OUTDIR, exist_ok=True)
    out = os.path.join(OUTDIR, f'轉倉位追蹤_{month}.xlsx')
    wb.save(out)
    print('saved', out, f'{n} days, {len(strikes)} strikes')


if __name__ == '__main__':
    month, s, e = sys.argv[1], sys.argv[2], sys.argv[3]
    y, m, d = map(int, s.split('-')); s = date(y, m, d)
    y, m, d = map(int, e.split('-')); e = date(y, m, d)
    build(month, s, e)
